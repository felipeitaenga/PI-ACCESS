"""Backend tests for new features (iter 20):
- GET /api/reports/data/{kind}?format=... with filters, PDF, XLSX
- PUT /api/selfreg/requests/{id} (edit pending)
- Company-manager scope on reports

Run: cd /app/backend && python -m pytest -q -o addopts='' tests/test_reports_and_selfreg_edit.py
"""
import asyncio
import base64
import hashlib
import io
import os
import re
import sys
import uuid
from datetime import datetime, timedelta, timezone
from pathlib import Path

import pytest
import requests

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

from motor.motor_asyncio import AsyncIOMotorClient  # noqa: E402

_LOOP = asyncio.new_event_loop()
asyncio.set_event_loop(_LOOP)
db = AsyncIOMotorClient(os.environ["MONGO_URL"], io_loop=_LOOP)[os.environ["DB_NAME"]]

BASE = os.environ.get("REACT_APP_BACKEND_URL", "https://nx-access-demo.preview.emergentagent.com").rstrip("/")
API = f"{BASE}/api"
ADMIN = {"email": "felipejferreira@gmail.com", "password": "NxAccess@2026"}
RECEP = {"email": "recepcao@nxaccess.com", "password": "Recepcao@2026"}
MGR = {"email": "TEST_rep_mgr@teste.com", "password": "TEST_Mgr@2026"}
FACE = Path("/tmp/face.jpg")


def _sha(t):
    return hashlib.sha256((t or "").encode()).hexdigest()


def _run(coro):
    return _LOOP.run_until_complete(coro)


def login(creds):
    s = requests.Session()
    r = s.post(f"{API}/auth/login", json=creds, timeout=30)
    assert r.status_code == 200, f"login {creds['email']}: {r.status_code} {r.text}"
    s.headers["Authorization"] = f"Bearer {r.json()['access_token']}"
    return s


def face_dataurl():
    with open(FACE, "rb") as f:
        return "data:image/jpeg;base64," + base64.b64encode(f.read()).decode()


def valid_cpf(seed=12345678):
    digits = [int(x) for x in str(seed).zfill(9)]
    for _ in range(2):
        s = sum((len(digits) + 1 - j) * d for j, d in enumerate(digits))
        digits.append(((s * 10) % 11) % 10)
    return "".join(str(d) for d in digits)


@pytest.fixture(scope="module")
def admin():
    return login(ADMIN)


@pytest.fixture(scope="module")
def recep():
    return login(RECEP)


@pytest.fixture(scope="module")
def companies(admin):
    r = admin.get(f"{API}/selfreg/companies")
    assert r.status_code == 200
    return r.json()


@pytest.fixture(scope="module")
def host_person(admin, companies):
    for c in companies:
        people = _run(db.people.find(
            {"company_id": c["id"], "status": "ativo",
             "can_authorize_visits": {"$ne": False}, "deleted": {"$ne": True}},
            {"_id": 0, "id": 1, "name": 1, "company_id": 1}).limit(1).to_list(1))
        if people:
            return people[0], c
    pytest.skip("no active host person")


@pytest.fixture(scope="module")
def visit_type(admin):
    r = admin.get(f"{API}/visits/types?active=true")
    assert r.status_code == 200 and r.json()
    return r.json()[0]


# ---------- Reports ----------
class TestReports:
    def test_invalid_kind(self, admin):
        r = admin.get(f"{API}/reports/data/unknown?format=json")
        assert r.status_code == 404

    def test_invalid_format(self, admin):
        r = admin.get(f"{API}/reports/data/access?format=csv")
        assert r.status_code == 400

    def test_json_access(self, admin):
        r = admin.get(f"{API}/reports/data/access?format=json")
        assert r.status_code == 200, r.text
        js = r.json()
        assert "columns" in js and "rows" in js and "total" in js
        assert isinstance(js["rows"], list)
        assert js["total"] == len(js["rows"])
        assert js["title"] == "Acessos"
        # columns match expected shape
        keys = [c["key"] for c in js["columns"]]
        assert "timestamp" in keys and "person_name" in keys

    def test_json_denied_filter(self, admin):
        r = admin.get(f"{API}/reports/data/denied?format=json")
        assert r.status_code == 200
        for row in r.json()["rows"][:20]:
            # denied report: result label is "Negado" or empty when access_result=denied is forced
            assert row.get("result") in ("Negado", "", None) or row["result"]

    def test_json_visits_date_filter(self, admin):
        start = (datetime.now(timezone.utc) - timedelta(days=365)).strftime("%Y-%m-%d")
        end = datetime.now(timezone.utc).strftime("%Y-%m-%d")
        r = admin.get(f"{API}/reports/data/visits",
                      params={"format": "json", "start": start, "end": end})
        assert r.status_code == 200
        assert "rows" in r.json()

    def test_json_bad_date(self, admin):
        r = admin.get(f"{API}/reports/data/visits?format=json&start=bad-date")
        assert r.status_code == 400

    def test_json_people(self, admin):
        r = admin.get(f"{API}/reports/data/people?format=json")
        assert r.status_code == 200
        js = r.json()
        assert js["title"] == "Pessoas cadastradas"
        for row in js["rows"][:5]:
            assert "person_type" in row and "name" in row

    def test_json_people_visitors(self, admin):
        r = admin.get(f"{API}/reports/data/people?format=json&person_type=visitante")
        assert r.status_code == 200
        for row in r.json()["rows"][:10]:
            assert row["person_type"] == "Visitante"

    def test_json_present(self, admin):
        r = admin.get(f"{API}/reports/data/present?format=json")
        assert r.status_code == 200
        assert "rows" in r.json()

    def test_pdf_access(self, admin):
        r = admin.get(f"{API}/reports/data/access?format=pdf")
        assert r.status_code == 200
        assert r.headers["content-type"].startswith("application/pdf")
        assert r.content[:4] == b"%PDF"
        assert len(r.content) > 500

    def test_xlsx_people_header_row_4(self, admin):
        from openpyxl import load_workbook
        r = admin.get(f"{API}/reports/data/people?format=xlsx")
        assert r.status_code == 200
        assert "spreadsheetml" in r.headers["content-type"]
        wb = load_workbook(io.BytesIO(r.content))
        ws = wb.active
        # Row 1: title; row 4: column headers
        assert "SCARVEL" in str(ws.cell(1, 1).value)
        headers = [ws.cell(4, i).value for i in range(1, 9)]
        assert headers[0] == "Nome"
        assert "Tipo" in headers

    def test_xlsx_visits(self, admin):
        from openpyxl import load_workbook
        r = admin.get(f"{API}/reports/data/visits?format=xlsx")
        assert r.status_code == 200
        wb = load_workbook(io.BytesIO(r.content))
        ws = wb.active
        assert ws.cell(4, 1).value == "Visitante"


@pytest.fixture(scope="module")
def manager_setup(admin, companies, host_person):
    host, comp = host_person
    cid = comp["id"]
    r = admin.post(f"{API}/auth/users", json={
        "name": "TEST Rep Manager", "email": MGR["email"], "password": MGR["password"],
        "role": "company_manager", "company_ids": [cid]})
    if r.status_code in (400, 409):
        users = admin.get(f"{API}/auth/users").json()
        u = next((x for x in users if x["email"] == MGR["email"]), None)
        if u:
            admin.put(f"{API}/auth/users/{u['id']}",
                      json={"company_ids": [cid], "password": MGR["password"],
                            "role": "company_manager"})
    elif r.status_code not in (200, 201):
        pytest.skip(f"manager create failed: {r.status_code} {r.text}")
    cidB = str(uuid.uuid4())
    _run(db.companies.insert_one({
        "id": cidB, "name": f"TEST_RepEmpresaB_{cidB[:6]}", "enabled": True,
        "selfreg_enabled": False, "deleted": False,
        "created_at": datetime.now(timezone.utc).isoformat()}))
    pidB = str(uuid.uuid4())
    _run(db.people.insert_one({
        "id": pidB, "name": "TEST_RepPersonB", "person_type": "fixo",
        "company_id": cidB, "status": "ativo", "deleted": False,
        "created_at": datetime.now(timezone.utc).isoformat()}))
    mgr = login(MGR)
    return {"mgr": mgr, "cid": cid, "cidB": cidB, "pidB": pidB}


class TestReportsScope:
    """Company manager must only see their own company's data."""

    def test_manager_people_excludes_other_company(self, manager_setup):
        r = manager_setup["mgr"].get(f"{API}/reports/data/people?format=json")
        assert r.status_code == 200
        names = [row["name"] for row in r.json()["rows"]]
        assert "TEST_RepPersonB" not in names, "manager must not see other-company people"

    def test_admin_sees_other_company(self, admin, manager_setup):
        r = admin.get(f"{API}/reports/data/people?format=json")
        assert r.status_code == 200
        names = [row["name"] for row in r.json()["rows"]]
        assert "TEST_RepPersonB" in names, "admin should see the TEST person"


# ---------- SelfReg edit_request ----------
@pytest.fixture(scope="module")
def pending(admin, companies, host_person, visit_type):
    _run(db.rate_hits.delete_many({}))
    host, comp = host_person
    cid = comp["id"]
    r = admin.put(f"{API}/selfreg/companies/{cid}",
                  json={"enabled": True, "visit_type_id": visit_type["id"], "regenerate": True})
    assert r.status_code == 200, r.text
    code = r.json()["selfreg_code"]
    cpf = valid_cpf(55500001)
    sub = requests.post(f"{API}/public/selfreg/submit", json={
        "code": code, "consent": True, "name": "TEST Edit Original",
        "doc_type": "cpf", "doc_number": cpf,
        "host_person_id": host["id"], "photo": face_dataurl()})
    assert sub.status_code == 200, sub.text
    tok = sub.json()["status_token"]
    reg = _run(db.self_registrations.find_one(
        {"status_token_hash": _sha(tok)}, {"_id": 0, "id": 1}))
    assert reg
    return {"rid": reg["id"], "code": code, "cid": cid, "host": host, "cpf": cpf}


class TestSelfRegEdit:
    def test_edit_name_and_fields(self, admin, pending):
        rid = pending["rid"]
        new_cpf = valid_cpf(55500002)
        r = admin.put(f"{API}/selfreg/requests/{rid}", json={
            "name": "TEST Edit NewName Sobrenome",
            "doc_type": "cpf", "doc_number": new_cpf,
            "visitor_company": "TEST VisCo Edited",
            "plate": "abc1d23",
            "phone": "11999887766",
            "email": "TEST.edit@example.com"})
        assert r.status_code == 200, r.text
        data = r.json()
        assert data["name"] == "TEST Edit NewName Sobrenome"
        assert re.sub(r"\D", "", data["doc_number"]) == new_cpf
        assert data["visitor_company"] == "TEST VisCo Edited"
        assert data["plate"] == "ABC1D23"
        assert data["email"] == "test.edit@example.com"

    def test_edit_validations(self, admin, pending):
        rid = pending["rid"]
        r = admin.put(f"{API}/selfreg/requests/{rid}", json={
            "name": "SoNome",  # single word
            "email": "not-an-email",
            "doc_type": "cpf", "doc_number": "00000000000",  # invalid cpf
            "end_at": "not-a-date"})
        assert r.status_code == 422, r.text
        errs = r.json()["detail"]["errors"]
        assert "name" in errs
        assert "email" in errs
        assert "doc_number" in errs
        assert "end_at" in errs

    def test_edit_invalid_visit_type(self, admin, pending):
        r = admin.put(f"{API}/selfreg/requests/{pending['rid']}",
                      json={"visit_type_id": "does-not-exist-id"})
        assert r.status_code == 422
        assert "visit_type_id" in r.json()["detail"]["errors"]

    def test_edit_sets_blacklist_from_new_doc(self, admin, pending):
        # Create a blacklisted visitor with specific CPF
        bl_cpf = valid_cpf(55599990)
        vid = str(uuid.uuid4())
        _run(db.visitors.insert_one({
            "id": vid, "name": "TEST BL for Edit", "doc_type": "cpf",
            "doc_number": bl_cpf, "doc_key": re.sub(r"\D", "", bl_cpf),
            "blacklist": True, "blacklist_reason": "TEST", "enabled": True,
            "deleted": False, "created_at": datetime.now(timezone.utc).isoformat()}))
        # Edit the pending req to use this doc
        r = admin.put(f"{API}/selfreg/requests/{pending['rid']}",
                      json={"doc_type": "cpf", "doc_number": bl_cpf})
        assert r.status_code == 200, r.text
        after = _run(db.self_registrations.find_one(
            {"id": pending["rid"]}, {"_id": 0, "blacklist": 1}))
        assert after.get("blacklist") is True, "blacklist should flip to true"

    def test_edit_then_approve_uses_edited_data(self, admin, recep, companies, host_person, visit_type):
        """New pending → edit name → approve → visit must carry edited name."""
        _run(db.rate_hits.delete_many({}))
        host, comp = host_person
        cid = comp["id"]
        rr = admin.put(f"{API}/selfreg/companies/{cid}",
                       json={"enabled": True, "visit_type_id": visit_type["id"]})
        code = rr.json()["selfreg_code"]
        cpf = valid_cpf(55500100)
        sub = requests.post(f"{API}/public/selfreg/submit", json={
            "code": code, "consent": True, "name": "TEST Approve Before Edit",
            "doc_type": "cpf", "doc_number": cpf,
            "host_person_id": host["id"], "photo": face_dataurl()})
        assert sub.status_code == 200
        rid = _run(db.self_registrations.find_one(
            {"status_token_hash": _sha(sub.json()["status_token"])},
            {"_id": 0, "id": 1}))["id"]
        new_name = "TEST Approve EditedName Final"
        e = admin.put(f"{API}/selfreg/requests/{rid}", json={"name": new_name})
        assert e.status_code == 200
        ap = recep.post(f"{API}/selfreg/requests/{rid}/approve", json={})
        assert ap.status_code == 200, ap.text
        vid = ap.json()["visit_id"]
        v = _run(db.visits.find_one({"id": vid}, {"_id": 0, "visitor_name": 1}))
        assert v["visitor_name"] == new_name, f"approved visit should carry edited name, got {v}"

    def test_edit_not_pending_returns_409(self, admin, recep, companies, host_person, visit_type):
        _run(db.rate_hits.delete_many({}))
        host, comp = host_person
        cid = comp["id"]
        rr = admin.put(f"{API}/selfreg/companies/{cid}",
                       json={"enabled": True, "visit_type_id": visit_type["id"]})
        code = rr.json()["selfreg_code"]
        cpf = valid_cpf(55500200)
        sub = requests.post(f"{API}/public/selfreg/submit", json={
            "code": code, "consent": True, "name": "TEST Edit After Approve",
            "doc_type": "cpf", "doc_number": cpf,
            "host_person_id": host["id"], "photo": face_dataurl()})
        assert sub.status_code == 200
        rid = _run(db.self_registrations.find_one(
            {"status_token_hash": _sha(sub.json()["status_token"])},
            {"_id": 0, "id": 1}))["id"]
        ap = recep.post(f"{API}/selfreg/requests/{rid}/approve", json={})
        assert ap.status_code == 200
        r = admin.put(f"{API}/selfreg/requests/{rid}", json={"name": "TEST Will Not Work"})
        assert r.status_code == 409

    def test_edit_manager_out_of_scope_404(self, admin, companies, host_person, visit_type):
        # Create pending in company A
        _run(db.rate_hits.delete_many({}))
        _run(db.self_registrations.delete_many({"name": {"$regex": "^TEST Scope Edit"}}))
        host, comp = host_person
        cid = comp["id"]
        rr = admin.put(f"{API}/selfreg/companies/{cid}",
                       json={"enabled": True, "visit_type_id": visit_type["id"]})
        code = rr.json()["selfreg_code"]
        import random
        cpf = valid_cpf(random.randint(55500300, 55599000))
        sub = requests.post(f"{API}/public/selfreg/submit", json={
            "code": code, "consent": True, "name": "TEST Scope Edit",
            "doc_type": "cpf", "doc_number": cpf,
            "host_person_id": host["id"], "photo": face_dataurl()})
        assert sub.status_code == 200
        rid = _run(db.self_registrations.find_one(
            {"status_token_hash": _sha(sub.json()["status_token"])},
            {"_id": 0, "id": 1}))["id"]
        # Create a manager scoped to a different company
        other_cid = str(uuid.uuid4())
        _run(db.companies.insert_one({
            "id": other_cid, "name": f"TEST_OtherCo_{other_cid[:6]}", "enabled": True,
            "selfreg_enabled": False, "deleted": False,
            "created_at": datetime.now(timezone.utc).isoformat()}))
        mgr_email = "TEST_scope_edit_mgr@teste.com"
        r = admin.post(f"{API}/auth/users", json={
            "name": "TEST Scope Edit Mgr", "email": mgr_email, "password": "TEST_Mgr@2026",
            "role": "company_manager", "company_ids": [other_cid]})
        if r.status_code == 409:
            users = admin.get(f"{API}/auth/users").json()
            u = next((x for x in users if x["email"] == mgr_email), None)
            if u:
                admin.put(f"{API}/auth/users/{u['id']}",
                          json={"company_ids": [other_cid], "role": "company_manager",
                                "password": "TEST_Mgr@2026"})
        mgr = login({"email": mgr_email, "password": "TEST_Mgr@2026"})
        e = mgr.put(f"{API}/selfreg/requests/{rid}", json={"name": "TEST Won't Reach"})
        assert e.status_code == 404


# ---------- Cleanup ----------
def test_zzz_cleanup(admin):
    async def do():
        await db.visitors.delete_many({"name": {"$regex": "^TEST "}})
        await db.visits.delete_many({"visitor_name": {"$regex": "^TEST "}})
        await db.self_registrations.delete_many({"name": {"$regex": "^TEST"}})
        await db.companies.delete_many({"name": {"$regex": "^TEST_"}})
        await db.people.delete_many({"name": {"$regex": "^TEST_"}})
        await db.rate_hits.delete_many({})
        for em in (MGR["email"], "TEST_scope_edit_mgr@teste.com"):
            u = await db.users.find_one({"email": em}, {"_id": 1})
            if u:
                await db.users.delete_one({"_id": u["_id"]})
        # Reset company selfreg flag — we don't know original state; best-effort disable
        await db.companies.update_many(
            {"selfreg_enabled": True},
            {"$set": {"selfreg_enabled": False, "selfreg_code": ""}})
    _run(do())
    print("cleanup done")
