"""Relatórios com filtros e exportação (JSON, Excel, PDF)."""
import io
import re
from datetime import datetime, timedelta, timezone

from fastapi import APIRouter, Depends, HTTPException
from fastapi.responses import StreamingResponse

from database import db
from registry_common import ND
from scope import scope_filter, own_company_or_404
from security import get_current_user
from visits_core import TZ

router = APIRouter(prefix="/reports", tags=["reports"])
LIMIT = 5000
RESULT = {"granted": "Liberado", "denied": "Negado"}
VSTATUS = {"agendado": "Agendado", "presente": "Presente", "finalizado": "Finalizado", "expirado": "Expirado",
           "cancelado": "Cancelado"}

EV_COLS = [("timestamp", "Data/hora"), ("person_name", "Pessoa"), ("person_type", "Tipo"), ("company_name", "Empresa"),
           ("door_name", "Porta"), ("direction", "Sentido"), ("result", "Resultado"), ("reason", "Motivo"),
           ("credential_type", "Credencial")]
COLS = {
    "access": EV_COLS,
    "denied": EV_COLS,
    "visits": [("visitor_name", "Visitante"), ("doc", "Documento"), ("visitor_company", "Empresa do visitante"),
               ("host_name", "Anfitrião"), ("host_company_name", "Empresa visitada"), ("visit_type_name", "Tipo"),
               ("start_at", "Início"), ("end_at", "Fim previsto"), ("checkin_at", "Entrada"), ("checkout_at", "Saída"),
               ("status", "Status")],
    "present": [("name", "Pessoa"), ("kind", "Tipo"), ("company", "Empresa"), ("where", "Anfitrião / área"),
                ("since", "Desde"), ("until", "Até")],
    "people": [("name", "Nome"), ("person_type", "Tipo"), ("doc", "Documento"), ("company_name", "Empresa"),
               ("department", "Departamento"), ("categories", "Categorias"), ("status", "Status"),
               ("valid_until", "Validade")],
}
TITLES = {"access": "Acessos", "denied": "Negados e desconhecidos", "visits": "Visitas", "present": "Presentes agora",
          "people": "Pessoas cadastradas"}


def fmt(v):
    if not v:
        return ""
    try:
        d = datetime.fromisoformat(str(v).replace("Z", "+00:00"))
        return (d if d.tzinfo else d.replace(tzinfo=timezone.utc)).astimezone(TZ).strftime("%d/%m/%Y %H:%M")
    except ValueError:
        return str(v)


def day_bounds(start, end):
    out = {}
    try:
        if start:
            out["$gte"] = datetime.fromisoformat(start[:10]).replace(tzinfo=TZ).astimezone(timezone.utc).isoformat()
        if end:
            d = datetime.fromisoformat(end[:10]).replace(tzinfo=TZ) + timedelta(days=1)
            out["$lt"] = d.astimezone(timezone.utc).isoformat()
    except ValueError:
        raise HTTPException(status_code=400, detail="Período inválido")
    return out


def name_rx(s):
    return {"$regex": re.escape(s.strip()[:80]), "$options": "i"}


def _ptype(t):
    return "Visitante" if t == "visitante" else ("Desconhecido" if not t else t.capitalize())


async def _events(user, f, denied):
    q = dict(await scope_filter(user, "access_events"))
    rng = day_bounds(f.get("start"), f.get("end"))
    if rng:
        q["timestamp"] = rng
    if f.get("company_id"):
        q["company_id"] = f["company_id"]
    if f.get("person"):
        q["person_name"] = name_rx(f["person"])
    if f.get("person_type") == "visitante":
        q["person_type"] = "visitante"
    elif f.get("person_type") == "fixo":
        q["person_type"] = {"$nin": ["visitante", "", None]}
    if f.get("door_id"):
        q["door_id"] = f["door_id"]
    if denied:
        q["access_result"] = "denied"
        if f.get("denied_kind") == "desconhecido":
            q["$or"] = [{"stranger": True}, {"person_id": {"$in": [None, ""]}}]
        elif f.get("denied_kind") == "negado":
            q["stranger"] = {"$ne": True}
            q["person_id"] = {"$nin": [None, ""]}
    elif f.get("result") in RESULT:
        q["access_result"] = f["result"]
    if f.get("category_id"):
        cat = await db.categories.find_one({"id": f["category_id"]}, {"_id": 0, "code": 1})
        q["category_code"] = (cat or {}).get("code", "-")
    rows = []
    async for e in db.access_events.find(q, {"_id": 0, "captured_photo": 0, "person_photo": 0}).sort("timestamp", -1).limit(LIMIT):
        rows.append({"id": e.get("id"), "timestamp": fmt(e.get("timestamp")), "person_name": e.get("person_name") or "Desconhecido",
                     "person_type": _ptype(e.get("person_type")), "company_name": e.get("company_name") or "",
                     "door_name": e.get("door_name") or "", "direction": e.get("direction") or "",
                     "result": RESULT.get(e.get("access_result"), e.get("access_result") or ""),
                     "reason": e.get("reason") or "", "credential_type": e.get("credential_type") or ""})
    return rows


async def _visits(user, f):
    q = {**ND, **(await scope_filter(user, "visits"))}
    rng = day_bounds(f.get("start"), f.get("end"))
    if rng:
        q["start_at"] = rng
    if f.get("company_id"):
        q["host_company_id"] = f["company_id"]
    if f.get("person"):
        q["$or"] = [{"visitor_name": name_rx(f["person"])}, {"host_name": name_rx(f["person"])}]
    if f.get("category_id"):
        q["category_ids"] = f["category_id"]
    if f.get("status") in VSTATUS:
        q["status"] = f["status"]
    rows = []
    async for v in db.visits.find(q, {"_id": 0}).sort("start_at", -1).limit(LIMIT):
        rows.append({"id": v["id"], "visitor_name": v.get("visitor_name", ""), "doc": v.get("doc_number") or v.get("visitor_document", ""),
                     "visitor_company": v.get("visitor_company") or v.get("origin_company", ""), "host_name": v.get("host_name", ""),
                     "host_company_name": v.get("host_company_name", ""), "visit_type_name": v.get("visit_type_name", ""),
                     "start_at": fmt(v.get("start_at") or v.get("expected_start")), "end_at": fmt(v.get("end_at") or v.get("expected_end")),
                     "checkin_at": fmt(v.get("checkin_at")), "checkout_at": fmt(v.get("checkout_at")),
                     "status": VSTATUS.get(v.get("status"), v.get("status", ""))})
    return rows


async def _present(user, f):
    rows = []
    if f.get("person_type") != "fixo":
        q = {**ND, **(await scope_filter(user, "visits")), "status": "presente"}
        if f.get("company_id"):
            q["host_company_id"] = f["company_id"]
        if f.get("person"):
            q["visitor_name"] = name_rx(f["person"])
        async for v in db.visits.find(q, {"_id": 0}).sort("checkin_at", -1).limit(LIMIT):
            rows.append({"id": v["id"], "name": v.get("visitor_name", ""), "kind": "Visitante",
                         "company": v.get("host_company_name", ""), "where": v.get("host_name", ""),
                         "since": fmt(v.get("checkin_at") or v.get("start_at")), "until": fmt(v.get("end_at") or v.get("expected_end"))})
    if f.get("person_type") != "visitante":
        q = {**ND, **(await scope_filter(user, "people")), "person_type": {"$ne": "visitante"},
             "current_area_id": {"$nin": [None, ""]}}
        if f.get("company_id"):
            q["company_id"] = f["company_id"]
        if f.get("person"):
            q["name"] = name_rx(f["person"])
        if f.get("category_id"):
            q["category_ids"] = f["category_id"]
        areas = {a["id"]: a for a in await db.areas.find({}, {"_id": 0, "id": 1, "description": 1, "is_exterior": 1}).to_list(500)}
        async for p in db.people.find(q, {"_id": 0, "photo": 0, "photo_thumb": 0}).sort("name", 1).limit(LIMIT):
            a = areas.get(p["current_area_id"]) or {}
            if a.get("is_exterior"):
                continue
            rows.append({"id": p["id"], "name": p.get("name", ""), "kind": _ptype(p.get("person_type")),
                         "company": p.get("company_name", ""), "where": a.get("description", ""), "since": "", "until": ""})
    return rows


async def _people(user, f):
    cats = {c["id"]: c.get("code") or c.get("description", "") for c in
            await db.categories.find({}, {"_id": 0, "id": 1, "code": 1, "description": 1}).to_list(500)}
    rows = []
    if f.get("person_type") == "visitante":
        q = {**ND, **(await scope_filter(user, "visitors"))}
        if f.get("person"):
            q["name"] = name_rx(f["person"])
        async for v in db.visitors.find(q, {"_id": 0, "photo": 0, "photo_thumb": 0}).sort("name", 1).limit(LIMIT):
            rows.append({"id": v["id"], "name": v.get("name", ""), "person_type": "Visitante",
                         "doc": v.get("doc_number") or v.get("document", ""), "company_name": v.get("origin_company", ""),
                         "department": "", "categories": "", "status": "Lista negra" if v.get("blacklist") else "Ativo",
                         "valid_until": ""})
        return rows
    q = {**ND, **(await scope_filter(user, "people")), "person_type": {"$ne": "visitante"}}
    if f.get("company_id"):
        q["company_id"] = f["company_id"]
    if f.get("person"):
        q["name"] = name_rx(f["person"])
    if f.get("category_id"):
        q["category_ids"] = f["category_id"]
    if f.get("status"):
        q["status"] = f["status"]
    async for p in db.people.find(q, {"_id": 0, "photo": 0, "photo_thumb": 0}).sort("name", 1).limit(LIMIT):
        rows.append({"id": p["id"], "name": p.get("name", ""), "person_type": _ptype(p.get("person_type")),
                     "doc": p.get("doc_number") or p.get("cpf") or p.get("document", ""), "company_name": p.get("company_name", ""),
                     "department": p.get("department", ""), "categories": ", ".join(cats.get(c, "") for c in p.get("category_ids") or []),
                     "status": (p.get("status") or "").capitalize(), "valid_until": fmt(p.get("valid_until"))[:10]})
    return rows


async def build(kind, user, f):
    # gestor só filtra pelas próprias empresas (o filtro substituiria a restrição de escopo)
    own_company_or_404(user, f.get("company_id"))
    if kind in ("access", "denied"):
        return await _events(user, f, kind == "denied")
    if kind == "visits":
        return await _visits(user, f)
    if kind == "present":
        return await _present(user, f)
    return await _people(user, f)


def to_xlsx(kind, rows, subtitle):
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill
    wb = Workbook()
    ws = wb.active
    ws.title = TITLES[kind][:30]
    ws.append([f"PI-ACCESS · {TITLES[kind]}"])
    ws["A1"].font = Font(bold=True, size=13)
    ws.append([subtitle])
    ws.append([])
    ws.append([label for _, label in COLS[kind]])
    for c in ws[4]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="E8590C")
    for r in rows:
        ws.append([r.get(k, "") for k, _ in COLS[kind]])
    for i, (k, label) in enumerate(COLS[kind], start=1):
        width = max([len(label)] + [len(str(r.get(k, ""))) for r in rows[:300]])
        ws.column_dimensions[ws.cell(row=4, column=i).column_letter].width = min(48, width + 2)
    ws.freeze_panes = "A5"
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def to_pdf(kind, rows, subtitle):
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4, landscape
    from reportlab.lib.styles import getSampleStyleSheet
    from reportlab.lib.units import mm
    from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle
    buf = io.BytesIO()
    doc = SimpleDocTemplate(buf, pagesize=landscape(A4), leftMargin=10 * mm, rightMargin=10 * mm, topMargin=10 * mm,
                            bottomMargin=10 * mm, title=TITLES[kind])
    st = getSampleStyleSheet()
    cell = st["BodyText"].clone("cell", fontSize=7, leading=8.5)
    head = cell.clone("head", textColor=colors.white, fontName="Helvetica-Bold")
    data = [[Paragraph(label, head) for _, label in COLS[kind]]]
    for r in rows:
        data.append([Paragraph(str(r.get(k, "")).replace("&", "&amp;").replace("<", "&lt;"), cell) for k, _ in COLS[kind]])
    t = Table(data, repeatRows=1, colWidths=[doc.width / len(COLS[kind])] * len(COLS[kind]))
    t.setStyle(TableStyle([("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#E8590C")),
                           ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#F4F4F5")]),
                           ("GRID", (0, 0), (-1, -1), 0.25, colors.HexColor("#D4D4D8")), ("VALIGN", (0, 0), (-1, -1), "TOP")]))
    story = [Paragraph(f"PI-ACCESS · {TITLES[kind]}", st["Title"]), Paragraph(subtitle, st["Normal"]), Spacer(1, 4 * mm), t]
    if not rows:
        story.append(Paragraph("Nenhum registro encontrado.", st["Normal"]))
    doc.build(story)
    return buf.getvalue()


@router.get("/data/{kind}")
async def report(kind: str, format: str = "json", start: str = "", end: str = "", company_id: str = "", person: str = "",
                 person_type: str = "", door_id: str = "", result: str = "", category_id: str = "", status: str = "",
                 denied_kind: str = "",
                 user: dict = Depends(get_current_user)):
    fmt_ = format
    if kind not in COLS:
        raise HTTPException(status_code=404, detail="Relatório não encontrado")
    f = {"start": start, "end": end, "company_id": company_id, "person": person, "person_type": person_type,
         "door_id": door_id, "result": result, "category_id": category_id, "status": status, "denied_kind": denied_kind}
    rows = await build(kind, user, f)
    if fmt_ == "json":
        return {"title": TITLES[kind], "columns": [{"key": k, "label": label} for k, label in COLS[kind]],
                "rows": rows, "total": len(rows), "limit": LIMIT}
    period = f"{start[:10] or 'início'} a {end[:10] or 'hoje'}" if kind in ("access", "denied", "visits") else "situação atual"
    subtitle = f"{period} · {len(rows)} registro(s) · gerado em {datetime.now(TZ).strftime('%d/%m/%Y %H:%M')}"
    stamp = datetime.now(TZ).strftime("%Y%m%d-%H%M")
    if fmt_ == "xlsx":
        data, mime, ext = to_xlsx(kind, rows, subtitle), "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", "xlsx"
    elif fmt_ == "pdf":
        data, mime, ext = to_pdf(kind, rows, subtitle), "application/pdf", "pdf"
    else:
        raise HTTPException(status_code=400, detail="Formato inválido")
    return StreamingResponse(io.BytesIO(data), media_type=mime,
                             headers={"Content-Disposition": f'attachment; filename="relatorio-{kind}-{stamp}.{ext}"'})
